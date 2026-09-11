__all__ = [
    'PPTXVisual',
]

import io
from collections.abc import Sequence
from typing import Any

import numpy as np
import numpy.typing as npt
import pptx
import pptx.dml.color
import pptx.enum.shapes
import pptx.enum.text
import pptx.parts.image
import pptx.presentation
import pptx.shapes.autoshape
import pptx.shapes.connector
import pptx.shapes.picture
import pptx.shapes.shapetree
import pptx.slide
import pptx.util
from PIL import Image
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

from todd.colors import RGB, Color

from ..registries import VisualRegistry
from .base import BaseVisual, Pen, Point, TextStyle


@VisualRegistry.register_()
class PPTXVisual(BaseVisual):
    """Visualize data in the format of PowerPoint.

    The PowerPoint contains only one slide.
    For more details, refer to python-pptx_.

    .. _python-pptx: https://github.com/scanny/python-pptx
    """

    def __init__(
        self,
        *args,
        width: int,
        height: int,
        **kwargs,
    ) -> None:
        """Initialize the PowerPoint with a single slide."""
        super().__init__(*args, width=width, height=height, **kwargs)
        self._presentation = pptx.Presentation()
        self.presentation.slide_width = pptx.util.Pt(width)
        self.presentation.slide_height = pptx.util.Pt(height)

        slides: pptx.slide.Slides = self.presentation.slides
        slides.add_slide(self.presentation.slide_layouts[6])

    @property
    def width(self) -> int:
        width = self.presentation.slide_width
        return int(width.pt)  # type: ignore[union-attr]

    @property
    def height(self) -> int:
        height = self.presentation.slide_height
        return int(height.pt)  # type: ignore[union-attr]

    @property
    def presentation(self) -> pptx.presentation.Presentation:
        return self._presentation

    @property
    def slide(self) -> pptx.slide.Slide:
        return self.presentation.slides[0]

    @property
    def shapes(self) -> pptx.shapes.shapetree.SlideShapes:
        return self.slide.shapes

    @staticmethod
    def _set_color_format(
        color_format: pptx.dml.color.ColorFormat,
        color: Color,
    ) -> None:
        red, green, blue, *_ = RGB.from_(color).to_tuple(normalized=False)
        color_format.rgb = pptx.dml.color.RGBColor(
            int(red),
            int(green),
            int(blue),
        )

    def save(self, path: Any) -> None:
        self.presentation.save(path)

    def point(
        self,
        point: Point,
        pen: Pen,
        *,
        shape_type: MSO_AUTO_SHAPE_TYPE = MSO_AUTO_SHAPE_TYPE.OVAL,
    ) -> pptx.shapes.autoshape.Shape:
        diameter = pen.width
        shape = self.shapes.add_shape(
            shape_type,
            pptx.util.Pt(point.x - diameter / 2),
            pptx.util.Pt(point.y - diameter / 2),
            pptx.util.Pt(diameter),
            pptx.util.Pt(diameter),
        )
        shape.fill.solid()
        self._set_color_format(shape.fill.fore_color, pen.color)
        shape.line.fill.background()
        return shape

    def line(
        self,
        start: Point,
        end: Point,
        pen: Pen,
    ) -> pptx.shapes.connector.Connector:
        connector = self.shapes.add_connector(
            pptx.enum.shapes.MSO_CONNECTOR.STRAIGHT,
            pptx.util.Pt(start.x),
            pptx.util.Pt(start.y),
            pptx.util.Pt(end.x),
            pptx.util.Pt(end.y),
        )
        connector.line.width = pptx.util.Pt(pen.width)
        self._set_color_format(connector.line.color, pen.color)
        return connector

    def fill(
        self,
        points: Sequence[Point],
        color: Color,
    ) -> pptx.shapes.autoshape.Shape:
        points = tuple(points)
        builder = self.shapes.build_freeform(
            points[0].x,
            points[0].y,
            pptx.util.Pt(1),
        )
        builder.add_line_segments(
            [(point.x, point.y) for point in points[1:]],
            close=True,
        )
        shape: pptx.shapes.autoshape.Shape = builder.convert_to_shape()
        shape.fill.solid()
        self._set_color_format(shape.fill.fore_color, color)
        shape.line.fill.background()
        return shape

    def text(
        self,
        text: str,
        position: Point,
        style: TextStyle,
        width: float | None = None,
        height: float | None = None,
    ) -> pptx.shapes.autoshape.Shape:
        if width is None:
            width = self.width - position.x
        if height is None:
            height = self.height - position.y
        textbox = self.shapes.add_textbox(
            pptx.util.Pt(position.x),
            pptx.util.Pt(position.y),
            pptx.util.Pt(width),
            pptx.util.Pt(height),
        )
        text_frame = textbox.text_frame
        text_frame.clear()
        text_frame.margin_left = 0
        text_frame.margin_top = 0
        text_frame.margin_right = 0
        text_frame.margin_bottom = 0
        text_frame.text = text
        text_frame.word_wrap = True
        text_frame.auto_size = pptx.enum.text.MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for paragraph in text_frame.paragraphs:
            paragraph.font.name = 'Times New Roman'
            paragraph.font.size = pptx.util.Pt(style.font_size)
            self._set_color_format(paragraph.font.color, style.color)
        return textbox

    def image(
        self,
        image: npt.NDArray[np.uint8],
        position: Point,
        width: float | None = None,
        height: float | None = None,
        opacity: float = 1,
    ) -> pptx.shapes.picture.Picture:
        assert 0 <= opacity <= 1

        height_, width_, channels = image.shape
        assert channels == 3
        width, height = self._get_image_wh(image, width, height)
        alpha = np.full(
            (height_, width_, 1),
            round(255 * opacity),
            dtype=np.uint8,
        )
        image_ = Image.fromarray(np.concatenate([image, alpha], axis=-1))
        with io.BytesIO() as f:
            image_.save(f, 'PNG')
            f.seek(0)
            picture = self.shapes.add_picture(
                f,
                pptx.util.Pt(position.x),
                pptx.util.Pt(position.y),
                pptx.util.Pt(width),
                pptx.util.Pt(height),
            )
        picture_image: pptx.parts.image.Image = picture.image
        assert picture_image.dpi == (72, 72)
        return picture

    def rectangle(
        self,
        left_top: Point,
        right_bottom: Point,
        fill: Color | None = None,
        pen: Pen | None = None,
    ) -> pptx.shapes.autoshape.Shape:
        left, top = left_top
        right, bottom = right_bottom
        rectangle = self.shapes.add_shape(
            pptx.enum.shapes.MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            pptx.util.Pt(left),
            pptx.util.Pt(top),
            pptx.util.Pt(right - left),
            pptx.util.Pt(bottom - top),
        )

        if pen is None:
            rectangle.line.fill.background()
        else:
            rectangle.line.width = pptx.util.Pt(pen.width)
            self._set_color_format(rectangle.line.color, pen.color)

        if fill is None:
            rectangle.fill.background()
        else:
            rectangle.fill.solid()
            self._set_color_format(rectangle.fill.fore_color, fill)
        return rectangle
